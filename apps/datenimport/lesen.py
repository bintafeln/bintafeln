"""Tabellen aus CSV- und Excel-Dateien lesen (ADR 0078), so tolerant wie möglich.

- CSV: Zeichenkodierung UTF-8 (mit/ohne BOM) oder Windows-1252, Trennzeichen ; , Tab oder |
  werden erkannt.
- Excel (.xlsx): jedes Blatt mit Inhalt ist eine eigene Tabelle. Datumszellen werden zu
  JJJJ-MM-TT, ganze Zahlen ohne „,0“.
- Die Kopfzeile muss nicht die erste Zeile sein: ``kopfzeile_finden()`` überspringt Titel- und
  Leerzeilen; auf der Zuordnungsseite lässt sie sich ändern.

Alles wird als Text zurückgegeben; die Prüfung der Werte macht das jeweilige Ziel.
"""

import csv
import datetime as dt
import io
from dataclasses import dataclass, field

MAX_ZEILEN = 50_000
TRENNER = [";", "\t", ",", "|"]


class LeseFehler(ValueError):
    pass


@dataclass
class Tabelle:
    name: str  # Datei, bei Excel „Datei › Blatt“
    zeilen: list[list[str]] = field(default_factory=list)
    hinweise: list[str] = field(default_factory=list)

    def als_dict(self) -> dict:
        return {"name": self.name, "zeilen": self.zeilen, "hinweise": self.hinweise}


def _text(roh: bytes) -> tuple[str, str]:
    for kodierung, name in (("utf-8-sig", "UTF-8"), ("cp1252", "Windows-1252")):
        try:
            return roh.decode(kodierung), name
        except UnicodeDecodeError:
            continue
    return roh.decode("latin-1"), "Latin-1"


def _trenner(text: str) -> str:
    """Das Zeichen, das in den ersten Zeilen am gleichmäßigsten und häufigsten vorkommt."""
    zeilen = [z for z in text.splitlines() if z.strip()][:10]
    if not zeilen:
        return ";"
    besser = None
    for t in TRENNER:
        anzahl = [z.count(t) for z in zeilen]
        if not max(anzahl):
            continue
        # gleiche Anzahl in vielen Zeilen schlägt einzelne Ausreißer (z. B. Komma im Titel)
        haeufigste = max(set(anzahl), key=anzahl.count)
        wert = (anzahl.count(haeufigste) * (haeufigste > 0), haeufigste)
        if besser is None or wert > besser[0]:
            besser = (wert, t)
    return besser[1] if besser else ";"


def csv_lesen(name: str, roh: bytes) -> Tabelle:
    text, kodierung = _text(roh)
    trenner = _trenner(text)
    tabelle = Tabelle(name=name)
    if kodierung != "UTF-8":
        tabelle.hinweise.append(f"Zeichenkodierung {kodierung} erkannt.")
    for werte in csv.reader(io.StringIO(text, newline=""), delimiter=trenner):
        tabelle.zeilen.append([w.strip() for w in werte])
        if len(tabelle.zeilen) > MAX_ZEILEN:
            raise LeseFehler(f"{name}: mehr als {MAX_ZEILEN} Zeilen.")
    return tabelle


def zelle(wert) -> str:
    """Excel-Zellwert als Text: Datum als JJJJ-MM-TT, 123.0 als 123, Wahrheitswert als ja/nein."""
    if wert is None:
        return ""
    if isinstance(wert, bool):
        return "ja" if wert else "nein"
    if isinstance(wert, dt.datetime):
        if wert.time() == dt.time(0, 0):
            return wert.date().isoformat()
        return wert.strftime("%Y-%m-%d %H:%M")
    if isinstance(wert, dt.date):
        return wert.isoformat()
    if isinstance(wert, dt.time):
        return wert.strftime("%H:%M")
    if isinstance(wert, float):
        return str(int(wert)) if wert.is_integer() else str(wert)
    return str(wert).strip()


def excel_lesen(name: str, roh: bytes) -> list[Tabelle]:
    from openpyxl import load_workbook

    try:
        mappe = load_workbook(io.BytesIO(roh), read_only=True, data_only=True)
    except Exception as fehler:  # beschädigte oder falsche Datei: verständliche Meldung
        raise LeseFehler(f"{name}: keine lesbare Excel-Datei (.xlsx).") from fehler
    tabellen = []
    try:
        for blatt in mappe.worksheets:
            zeilen = []
            for werte in blatt.iter_rows(values_only=True):
                zeilen.append([zelle(w) for w in werte])
                if len(zeilen) > MAX_ZEILEN:
                    raise LeseFehler(f"{name}: mehr als {MAX_ZEILEN} Zeilen.")
            while zeilen and not any(zeilen[-1]):
                zeilen.pop()
            if zeilen:
                titel = name if len(mappe.worksheets) == 1 else f"{name} › {blatt.title}"
                tabellen.append(Tabelle(name=titel, zeilen=zeilen))
    finally:
        mappe.close()
    return tabellen


def lesen(name: str, roh: bytes) -> list[Tabelle]:
    """Eine hochgeladene Datei als eine oder mehrere Tabellen."""
    endung = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if endung in ("xlsx", "xlsm"):
        tabellen = excel_lesen(name, roh)
    elif endung in ("xls", "ods", "numbers"):
        raise LeseFehler(f"{name}: Bitte in Excel als .xlsx oder als CSV speichern.")
    elif roh[:2] == b"PK":  # .xlsx ohne passende Endung
        tabellen = excel_lesen(name, roh)
    else:
        tabellen = [csv_lesen(name, roh)]
    tabellen = [t for t in tabellen if any(any(z) for z in t.zeilen)]
    if not tabellen:
        raise LeseFehler(f"{name}: Die Datei enthält keine Daten.")
    return tabellen


def kopfzeile_finden(zeilen: list[list[str]]) -> int:
    """Index der wahrscheinlichen Kopfzeile: die erste Zeile mit mindestens zwei gefüllten
    Zellen, die überwiegend Text (keine Zahlen oder Daten) enthalten."""
    breite = max((sum(1 for w in z if w) for z in zeilen[:30]), default=0)
    for i, z in enumerate(zeilen[:30]):
        gefuellt = [w for w in z if w]
        if len(gefuellt) < max(2, breite // 2):
            continue
        texte = [w for w in gefuellt if not any(c.isdigit() for c in w)]
        if len(texte) >= len(gefuellt) * 0.6:
            return i
    return 0
